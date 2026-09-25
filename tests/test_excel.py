import openpyxl
import pytest

import config as cfgmod
from excel_writer import invalid_filename_chars, normalize_filename, sanitize_sheet_name, validate_graphs, write_workbook
from helpers import EXPECTED_XLSX, SAMPLE_LOGGER, SAMPLE_UART, write_logger_csv, write_uart_csv
from logger_reader import read_logger_csv
from make_expected import make
from merger import build_analysis_table, merge
from models import ELAPSED_KEY, SOURCE_LOGGER, SOURCE_UART, GraphSetting, ItemSetting
from uart_reader import read_uart_csv


def _values(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_matches_expected_output(tmp_path):
    out = str(tmp_path / "out.xlsx")
    make(SAMPLE_UART, SAMPLE_LOGGER, out)
    got = openpyxl.load_workbook(out)
    exp = openpyxl.load_workbook(EXPECTED_XLSX)
    assert got.sheetnames == exp.sheetnames == ["解析", "グラフ", "log40", "260924-090155"]
    for name in ("解析", "log40", "260924-090155"):
        assert _values(got[name]) == _values(exp[name]), name
    assert len(got["グラフ"]._charts) == len(exp["グラフ"]._charts) == 1


def _setup(tmp_path, gaps=False):
    times = [i * 1000 for i in range(10)]
    if gaps:
        del times[3:5]
    u = read_uart_csv(write_uart_csv(str(tmp_path / "u.csv"), times))
    lg = read_logger_csv(write_logger_csv(str(tmp_path / "l.CSV"), 60, overrange_rows={10}))
    r = merge(u, lg, 1000, 200)
    return u, lg, r


def test_analysis_sheet_follows_settings(tmp_path):
    u, lg, r = _setup(tmp_path, gaps=True)
    items = [
        ItemSetting("t_ms", SOURCE_UART, False, "t_ms"),
        ItemSetting("voltage_mV", SOURCE_UART, True, "電圧(V)", 0.001, 0),
        ItemSetting("current_mA", SOURCE_UART, True, "電流(mA)", 1, 5),
        ItemSetting("temp_C", SOURCE_UART, False, "温度"),
        ItemSetting("CH1", SOURCE_LOGGER, False, "CH1(mV)", unit="mV"),
        ItemSetting("CH2", SOURCE_LOGGER, True, "CH2(mV)", 1000, 0, unit="V"),
    ]
    t = build_analysis_table(r, u, lg, items, "時間[s]")
    out = str(tmp_path / "o.xlsx")
    write_workbook(out, t, u, lg, [])
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["解析", "u", "l"]  # グラフ0個ならグラフシートなし
    ws = wb["解析"]
    assert ws["A1"].value == "測定日時"
    assert ws["B1"].value.strftime("%Y/%m/%d %H:%M:%S") == "2026/09/24 09:09:54"
    assert ws["B1"].number_format == "yyyy/mm/dd hh:mm:ss"
    assert (ws["A2"].value, ws["B2"].value) == ("測定機", "GL240")
    assert ws["A5"].value == "マイコン内部データ(UART)"
    assert [c.value for c in ws[6]] == ["時間[s]", "電圧(V)", "電流(mA)", None, "CH2(mV)"]
    assert ws["D5"].value is None and ws["E5"].value == "測定値"
    # 1行目
    assert ws["A7"].value == 0
    assert ws["B7"].value == pytest.approx(3.7)
    assert ws["C7"].value == -95
    assert ws["E7"].value == pytest.approx(2000)
    # 3行目（t=2s）→ ロガー行インデックス10 はレンジオーバー → 空欄
    assert ws["A9"].value == 2
    assert ws["E9"].value is None
    # 欠落行（4, 5行目）は UART 側が空欄で灰色、ロガーは値あり
    for row in (10, 11):
        assert ws[f"B{row}"].value is None and ws[f"C{row}"].value is None
        assert ws[f"E{row}"].value is not None
        assert ws[f"A{row}"].fill.fgColor.rgb.endswith("D9D9D9")
        assert ws[f"E{row}"].fill.fgColor.rgb.endswith("D9D9D9")
    assert ws["A12"].value == 5 and ws["B12"].value == pytest.approx(3.703)
    assert ws["A12"].fill.fill_type is None
    assert isinstance(ws["B7"].value, float)


def test_graph_sheet(tmp_path):
    u, lg, r = _setup(tmp_path)
    cfg = cfgmod.default_config()
    items, _ = cfgmod.build_items(cfg, u, lg)
    t = build_analysis_table(r, u, lg, items, "経過時間(s)")
    # グラフ1：横軸は初期値（elapsed_ms）、グラフ2：横軸は経過時間(s)
    graphs = [
        GraphSetting(["voltage_mV", "CH2"], ["current_mA", "CH1"]),  # 各軸2系列
        GraphSetting(["CH1"], [None], ELAPSED_KEY),
    ]
    {it.key: it for it in items}["CH2"].color = "FF0000"
    out = str(tmp_path / "o.xlsx")
    write_workbook(out, t, u, lg, graphs)

    import zipfile

    with zipfile.ZipFile(out) as z:
        # lxml の有無で "<a />" / "<a/>" の書き方が変わるのでそろえる
        c1 = z.read("xl/charts/chart1.xml").decode().replace(" />", "/>")
        c2 = z.read("xl/charts/chart2.xml").decode().replace(" />", "/>")
    last = 6 + len(r.rows)
    # 横軸 = 解析シートの elapsed_ms 列（B列）、横軸タイトルはそのラベル
    assert f"<xVal><numRef><f>'解析'!$B$7:$B${last}</f>" in c1
    assert "time(ms)" in c1
    # 第1軸 電圧 = C列、第2軸 電流 = D列（elapsed_ms が B列）
    assert f"'解析'!$C$7:$C${last}" in c1 and f"'解析'!$D$7:$D${last}" in c1
    assert c1.count("<scatterChart>") == 2  # 第2軸あり
    assert '<crosses val="max"/>' in c1
    assert "電圧(mV) / CH2(V)  |  電流(mA) / CH1(mV)" in c1  # グラフタイトル
    assert c1.count("<ser>") == 4
    # 第1軸 2系列目 = CH2（J列）、第2軸 2系列目 = CH1（I列）
    assert f"'解析'!$J$7:$J${last}" in c1 and f"'解析'!$I$7:$I${last}" in c1
    # 線の色は項目の色（CH2 は赤に変更、電圧は既定色）
    assert '<a:srgbClr val="FF0000"/>' in c1
    volt_color = {it.key: it for it in items}["voltage_mV"].color
    assert f'<a:srgbClr val="{volt_color}"/>' in c1
    # 系列番号は通し番号
    assert all(f'<idx val="{n}"/>' in c1 for n in range(4))
    assert '<dispBlanksAs val="gap"/>' in c1
    assert '<symbol val="none"/>' in c1
    # 文字が重ならない設定：タイトル・凡例・軸タイトルは overlay なし、目盛の数値は外側
    assert c1.count('<overlay val="0"/>') == 5  # グラフタイトル・横軸・第1軸・第2軸・凡例
    assert '<tickLblPos val="low"/>' in c1 and '<tickLblPos val="high"/>' in c1
    # グラフ2：第2軸なし、CH1 は UART の右に1列空けた位置（A + 6項目 + 空列 → I列）
    assert c2.count("<scatterChart>") == 1
    assert f"'解析'!$I$7:$I${last}" in c2
    assert f"<xVal><numRef><f>'解析'!$A$7:$A${last}</f>" in c2
    assert "経過時間(s)" in c2
    wb = openpyxl.load_workbook(out)
    assert len(wb["グラフ"]._charts) == 2


def test_validate_graphs():
    items = [ItemSetting("a", SOURCE_UART, True, "A"), ItemSetting("b", SOURCE_UART, False, "B")]
    assert validate_graphs([GraphSetting("a", None)], items)  # 横軸 elapsed_ms が採用されていない
    assert validate_graphs([GraphSetting("a", None, ELAPSED_KEY)], items) == []
    assert validate_graphs([GraphSetting("a", None, None)], items)
    assert validate_graphs([GraphSetting("a", None, "a")], items)  # 横軸と縦軸が同じ
    items.append(ItemSetting("elapsed_ms", SOURCE_UART, True, "time(ms)"))
    assert validate_graphs([GraphSetting("a", None)], items) == []
    assert validate_graphs([GraphSetting(None, None)], items)
    assert validate_graphs([GraphSetting("b", None)], items)
    assert validate_graphs([GraphSetting("a", "a")], items)
    assert validate_graphs([GraphSetting("a", "")], items)
    # 各軸2系列まで・同じ項目の重複はエラー
    items.append(ItemSetting("c", SOURCE_UART, True, "C"))
    items.append(ItemSetting("d", SOURCE_UART, True, "D"))
    assert validate_graphs([GraphSetting(["a", "c"], ["d", None])], items) == []
    assert validate_graphs([GraphSetting(["a", "a"], [None])], items)
    assert validate_graphs([GraphSetting(["a", "c"], ["c", None])], items)
    assert validate_graphs([GraphSetting(["a", ""], [None])], items)
    assert validate_graphs([GraphSetting(["a", "b"], [None])], items)  # b は採用なし


def test_sheet_and_file_names():
    used = {"解析", "グラフ"}
    assert sanitize_sheet_name("a/b:c*d?[e]", used) == "a_b_c_d__e_"
    assert len(sanitize_sheet_name("x" * 40, used)) == 31
    assert sanitize_sheet_name("解析", used) == "解析(2)"
    assert invalid_filename_chars('a<b>.xlsx') == ["<", ">"]
    assert normalize_filename("abc") == "abc.xlsx"
    assert normalize_filename("abc.XLSX") == "abc.XLSX"


def test_analysis_sheet_styles_and_tab_colors(tmp_path):
    u, lg, r = _setup(tmp_path, gaps=True)
    cfg = cfgmod.default_config()
    items, _ = cfgmod.build_items(cfg, u, lg)
    t = build_analysis_table(r, u, lg, items, "経過時間(s)")
    out = str(tmp_path / "o.xlsx")
    write_workbook(out, t, u, lg, [GraphSetting("voltage_mV", None)])
    wb = openpyxl.load_workbook(out)
    assert wb["解析"].sheet_properties.tabColor.rgb.endswith("A9D08E")
    assert wb["グラフ"].sheet_properties.tabColor.rgb.endswith("F4B084")
    assert wb["u"].sheet_properties.tabColor is None

    ws = wb["解析"]
    # A〜G: UART（経過時間 + 6項目）、H: 空列、I〜J: ロガー
    assert ws["A5"].font.b and ws["I5"].font.b
    # 見出し行：薄い緑。罫線は格子（細線）、ブロックの外周は中太線
    for col in "ABCDEFGIJ":
        c = ws[f"{col}6"]
        assert c.fill.fgColor.rgb.endswith("E2EFDA"), col
        assert c.border.top.style == "medium" and c.border.bottom.style == "thin", col
    assert ws["H6"].fill.fill_type is None and ws["H6"].border.left.style is None
    # データ部：全セルに細線の格子、外周（A列左・G列右・I列左・J列右・最終行下）は中太線
    assert ws["A7"].border.left.style == "medium" and ws["A7"].border.right.style == "thin"
    assert ws["C7"].border.left.style == "thin" and ws["C7"].border.top.style == "thin"
    assert ws["G7"].border.right.style == "medium"
    assert ws["I7"].border.left.style == "medium" and ws["J7"].border.right.style == "medium"
    assert ws["H7"].border.left.style is None and ws["H7"].border.bottom.style is None
    last = 6 + len(r.rows)
    for col in "ACGIJ":
        assert ws[f"{col}{last}"].border.bottom.style == "medium", col
    assert ws[f"C{last - 1}"].border.bottom.style == "thin"
    # 欠落行も格子付きで灰色
    assert ws["B10"].fill.fgColor.rgb.endswith("D9D9D9") and ws["B10"].border.left.style == "thin"


def test_time_unit_columns_for_graph_x_axis(tmp_path):
    u, lg, r = _setup(tmp_path, gaps=True)
    cfg = cfgmod.default_config()
    items, _ = cfgmod.build_items(cfg, u, lg)
    t = build_analysis_table(r, u, lg, items, "経過時間(s)")
    graphs = [
        GraphSetting(["voltage_mV"], [None], "elapsed_ms", "min"),  # time(ms) → time(min)
        GraphSetting(["current_mA"], [None], "elapsed_ms", "min"),  # 同じ単位は列 1 つ
        GraphSetting(["CH1"], [None], ELAPSED_KEY, "h"),  # 経過時間(s) → 経過時間(h)
        GraphSetting(["CH2"], [None], "elapsed_ms", "ms"),  # 元と同じ単位 → 列を足さない
    ]
    out = str(tmp_path / "o.xlsx")
    write_workbook(out, t, u, lg, graphs)
    ws = openpyxl.load_workbook(out)["解析"]
    heads = [c.value for c in ws[6]]
    # A 経過時間(s)、B 経過時間(h)、C time(ms)、D time(min)、以降は元どおり
    assert heads[:5] == ["経過時間(s)", "経過時間(h)", "time(ms)", "time(min)", "電圧(mV)"]
    assert heads.count("time(min)") == 1
    for row in range(7, 7 + len(r.rows)):
        ms = ws[f"C{row}"].value
        assert ws[f"D{row}"].value == pytest.approx(ms / 60000)
        assert ws[f"B{row}"].value == pytest.approx(ws[f"A{row}"].value / 3600, abs=1e-6)
    # 欠落行（10 行目）の換算値もある
    assert ws["D10"].value is not None and ws["D10"].fill.fgColor.rgb.endswith("D9D9D9")

    import zipfile

    last = 6 + len(r.rows)
    with zipfile.ZipFile(out) as z:
        c1 = z.read("xl/charts/chart1.xml").decode().replace(" />", "/>")
        c3 = z.read("xl/charts/chart3.xml").decode().replace(" />", "/>")
        c4 = z.read("xl/charts/chart4.xml").decode().replace(" />", "/>")
    assert f"<xVal><numRef><f>'解析'!$D$7:$D${last}</f>" in c1 and "time(min)" in c1
    assert f"<xVal><numRef><f>'解析'!$B$7:$B${last}</f>" in c3 and "経過時間(h)" in c3
    assert f"<xVal><numRef><f>'解析'!$C$7:$C${last}</f>" in c4


def test_time_unit_label():
    from excel_writer import time_unit_label

    assert time_unit_label("time(ms)", "min") == "time(min)"
    assert time_unit_label("経過時間(s)", "h") == "経過時間(h)"
    assert time_unit_label("time", "s") == "time(s)"


def test_effective_x_unit():
    assert GraphSetting(["a"], [None], "elapsed_ms", "min").effective_x_unit() == "min"
    assert GraphSetting(["a"], [None], "elapsed_ms", "ms").effective_x_unit() is None
    assert GraphSetting(["a"], [None], "voltage_mV", "min").effective_x_unit() is None  # 時間の項目でない
    assert GraphSetting(["a"], [None], ELAPSED_KEY, "min").effective_x_unit() == "min"
