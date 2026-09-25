"""Excel（.xlsx）の出力（仕様 6章）。"""

from __future__ import annotations

import math
import os
import re
from typing import Callable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import NumericAxis
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from logger_reader import LoggerData
from merger import AnalysisTable
from models import ELAPSED_KEY, GraphSetting, ItemSetting
from uart_reader import TIMESTAMP_COLUMN, UartData, parse_timestamp

SHEET_ANALYSIS = "解析"
SHEET_GRAPH = "グラフ"

HEADER_ROW = 6
DATA_START_ROW = 7

GAP_FILL = PatternFill("solid", fgColor="D9D9D9")
HEADER_FILL = PatternFill("solid", fgColor="E2EFDA")  # 見出し行（薄い緑）
TITLE_FONT = Font(bold=True, size=12)
THIN = Side(style="thin", color="000000")
# シート見出し（タブ）の色
TAB_COLOR_ANALYSIS = "A9D08E"  # 緑
TAB_COLOR_GRAPH = "F4B084"  # オレンジ

# グラフの大きさ（cm）と縦の間隔（行）
CHART_WIDTH_CM = 24
CHART_HEIGHT_CM = 12
CHART_ROW_STEP = 26

_INVALID_SHEET_CHARS = re.compile(r"[\\/?*\[\]:]")
_INVALID_FILE_CHARS = set('\\/:*?"<>|')

ProgressFn = Callable[[str], None]


def sanitize_sheet_name(name: str, used: set[str]) -> str:
    """Excel のシート名制限（31文字、\\ / ? * [ ] : 不可、重複不可）に合わせる。"""
    s = _INVALID_SHEET_CHARS.sub("_", name).strip("'") or "Sheet"
    s = s[:31]
    base = s
    n = 2
    while s.lower() in {u.lower() for u in used}:
        suffix = f"({n})"
        s = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(s)
    return s


def invalid_filename_chars(name: str) -> list[str]:
    return sorted({c for c in name if c in _INVALID_FILE_CHARS})


def normalize_filename(name: str) -> str:
    name = name.strip()
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    return name


def _raw_value(text: str):
    """CSV のセルを、Excel で開いたときと同じように数値化できれば数値にする。"""
    s = text.strip()
    if not s:
        return None
    t = s.replace(" ", "")
    if t[0] in "+-.0123456789" and "_" not in t:
        if t.isdigit():
            return int(t)
        try:
            v = float(t)
        except ValueError:
            return text
        if math.isfinite(v):
            return v
    return text


def _write_uart_raw(ws, uart: UartData) -> None:
    ts_idx = None
    for r, row in enumerate(uart.raw_rows):
        if r == 0:
            header = [h.strip() for h in row]
            ts_idx = header.index(TIMESTAMP_COLUMN) if TIMESTAMP_COLUMN in header else None
            ws.append(row)
            continue
        out = []
        for c, cell in enumerate(row):
            if c == ts_idx:
                ts = parse_timestamp(cell)
                if ts is not None:
                    wc = WriteOnlyCell(ws, value=ts)
                    wc.number_format = "yyyy/mm/dd hh:mm:ss.000"
                    out.append(wc)
                    continue
            out.append(_raw_value(cell))
        ws.append(out)


def _write_logger_raw(ws, logger: LoggerData, progress: ProgressFn | None) -> None:
    total = len(logger.raw_rows)
    for n, row in enumerate(logger.raw_rows):
        ws.append([_raw_value(c) for c in row])
        if progress and n % 50000 == 0 and n:
            progress(f"ロガーの元データを書き込み中… {n:,}/{total:,} 行")


def _write_analysis(ws, table: AnalysisTable, uart: UartData, logger: LoggerData, progress: ProgressFn | None) -> dict:
    """解析シートを書き、グラフ作成用に 項目キー -> 列番号 を返す。"""
    n_uart_cols = 1 + len(table.uart_items)
    logger_start = n_uart_cols + 2  # 1列空ける
    total_cols = n_uart_cols + (1 + len(table.logger_items) if table.logger_items else 0)

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 20
    for c in range(3, total_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12

    start = uart.timestamps[0].replace(microsecond=0)
    c_start = WriteOnlyCell(ws, value=start)
    c_start.number_format = "yyyy/mm/dd hh:mm:ss"
    ws.append(["測定日時", c_start])
    ws.append(["測定機", logger.model])
    ws.append([])
    ws.append([])

    uart_cols = range(1, n_uart_cols + 1)  # 1始まりの列番号
    logger_cols = range(logger_start, total_cols + 1) if table.logger_items else range(0)
    blocks = [uart_cols] + ([logger_cols] if table.logger_items else [])

    def styled(value, col: int, *, font=None, fill=None, border_all=False, bottom=False):
        c = WriteOnlyCell(ws, value=value)
        if font is not None:
            c.font = font
        if fill is not None:
            c.fill = fill
        if border_all:
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        else:
            # データ部は各ブロック（UART・ロガー）の外枠だけ線を引く
            left = right = None
            for b in blocks:
                if col == b.start:
                    left = THIN
                if col == b.stop - 1:
                    right = THIN
            if left or right or bottom:
                c.border = Border(left=left, right=right, bottom=THIN if bottom else None)
        return c

    row5 = [None] * total_cols
    row5[0] = styled("マイコン内部データ(UART)", 1, font=TITLE_FONT)
    if table.logger_items:
        row5[logger_start - 1] = styled("測定値", logger_start, font=TITLE_FONT)
    ws.append(row5)

    row6: list = [None] * total_cols
    labels6 = [table.elapsed_label] + [it.label for it in table.uart_items]
    for i, lab in enumerate(labels6):
        row6[i] = styled(lab, 1 + i, fill=HEADER_FILL, border_all=True)
    for i, it in enumerate(table.logger_items):
        row6[logger_start - 1 + i] = styled(it.label, logger_start + i, fill=HEADER_FILL, border_all=True)
    ws.append(row6)

    col_of: dict[str, int] = {}
    for i, it in enumerate(table.uart_items):
        col_of[it.key] = 2 + i
    for i, it in enumerate(table.logger_items):
        col_of.setdefault(it.key, logger_start + i)

    edge_cols = {c for b in blocks for c in (b.start, b.stop - 1)}
    in_block = set(uart_cols) | set(logger_cols)
    total = len(table.rows)
    for n, (sec, uv, lv, is_gap) in enumerate(table.rows):
        vals: list = [None] * total_cols
        vals[0] = sec
        for i, v in enumerate(uv):
            vals[1 + i] = v
        for i, v in enumerate(lv):
            vals[logger_start - 1 + i] = v
        is_last = n == total - 1
        for idx in range(total_cols):
            col = idx + 1
            if is_gap:
                vals[idx] = styled(vals[idx], col, fill=GAP_FILL, bottom=is_last and col in in_block)
            elif col in edge_cols or (is_last and col in in_block):
                vals[idx] = styled(vals[idx], col, bottom=is_last)
        ws.append(vals)
        if progress and n % 20000 == 0 and n:
            progress(f"解析シートを書き込み中… {n:,}/{total:,} 行")
    return col_of


def _series(ref_y, ref_x, label: str, color: str) -> Series:
    s = Series(ref_y, ref_x, title=label)
    s.marker.symbol = "none"
    s.smooth = False
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = 19050
    return s


def _add_charts(ws, ws_data, table: AnalysisTable, graphs: list[GraphSetting], col_of: dict[str, int]) -> None:
    """グラフシートに散布図を並べる。データは解析シート(ws_data)のセル範囲を参照する。"""
    first = DATA_START_ROW
    last = DATA_START_ROW + len(table.rows) - 1
    items = {it.key: it for it in table.uart_items + table.logger_items}
    labels = {k: it.label for k, it in items.items()}
    labels[ELAPSED_KEY] = table.elapsed_label
    col_of = {**col_of, ELAPSED_KEY: 1}

    def ref(col: int):
        return Reference(ws_data, min_col=col, min_row=first, max_row=last)

    for gi, g in enumerate(graphs):
        x_key = g.x or ELAPSED_KEY
        xref = ref(col_of[x_key])
        prim = [k for k in g.primary if k]
        sec = [k for k in g.secondary if k]

        c1 = ScatterChart()
        c1.scatterStyle = "lineMarker"
        c1.display_blanks = "gap"
        c1.width = CHART_WIDTH_CM
        c1.height = CHART_HEIGHT_CM
        for k in prim:
            c1.series.append(_series(ref(col_of[k]), xref, labels[k], items[k].color))
        p_title = " / ".join(labels[k] for k in prim)
        c1.x_axis.title = labels[x_key]
        c1.y_axis.title = p_title
        c1.x_axis.axPos = "b"
        c1.x_axis.delete = False
        c1.y_axis.delete = False
        title = p_title
        c2 = None
        if sec:
            s_title = " / ".join(labels[k] for k in sec)
            c2 = ScatterChart()
            c2.scatterStyle = "lineMarker"
            for k in sec:
                c2.series.append(_series(ref(col_of[k]), xref, labels[k], items[k].color))
            c2.y_axis.axId = 200
            c2.y_axis.title = s_title
            c2.y_axis.crosses = "max"
            c2.y_axis.axPos = "r"
            c2.y_axis.majorGridlines = None
            c2.y_axis.delete = False
            # 第2軸のグループには非表示の横軸を別に持たせる（Excel が作る2軸グラフと同じ構成）
            c2.x_axis = NumericAxis(axId=500, crossAx=200, delete=True, axPos="b")
            c2.x_axis.title = None
            c2.x_axis.majorGridlines = None
            c2.y_axis.crossAx = 500
            c1 += c2
            title = f"{p_title}  |  {s_title}"
        # 系列番号を通し番号にする（凡例・色の割り当てが重ならないように）
        for n, s_ in enumerate(c1.series + (c2.series if c2 else [])):
            s_.idx = s_.order = n
        c1.title = title
        c1.legend.position = "b"
        # タイトル・軸タイトル・凡例がグラフ本体や目盛の数値と重ならないよう、重ねない配置にする
        c1.title.overlay = False
        c1.legend.overlay = False
        c1.x_axis.title.overlay = False
        c1.y_axis.title.overlay = False
        # 目盛の数値は軸の外側（左端・下端・右端）に置き、軸タイトルとぶつからないようにする
        c1.x_axis.tickLblPos = "low"
        c1.x_axis.crosses = "min"  # 横軸は常にグラフの下端（縦軸に負の値があっても中央に来ない）
        c1.y_axis.crosses = "min"
        c1.y_axis.tickLblPos = "low"
        if c2 is not None:
            c2.y_axis.title.overlay = False
            c2.y_axis.tickLblPos = "high"
        anchor_row = 1 + gi * CHART_ROW_STEP
        ws.add_chart(c1, f"A{anchor_row}")


def validate_graphs(graphs: list[GraphSetting], items: list[ItemSetting]) -> list[str]:
    """グラフ設定の検査。エラーメッセージの一覧を返す（空なら OK）。"""
    enabled = {it.key for it in items if it.enabled}
    errors = []
    for i, g in enumerate(graphs, start=1):
        if not g.x or (g.x != ELAPSED_KEY and g.x not in enabled):
            errors.append(f"グラフ{i}：横軸を選択してください。")
        if not g.primary[0] or g.primary[0] not in enabled:
            errors.append(f"グラフ{i}：第1軸を選択してください。")
        for axis, keys in (("第1軸", g.primary[1:]), ("第2軸", g.secondary)):
            for k in keys:
                if k == "":
                    errors.append(f"グラフ{i}：{axis}の項目を選択してください（「なし」も選べます）。")
                elif k is not None and k not in enabled:
                    errors.append(f"グラフ{i}：{axis}の項目が採用されていません。")
        chosen = [k for k in [g.x] + g.series_keys() if k]
        if len(chosen) != len(set(chosen)):
            errors.append(f"グラフ{i}：横軸・縦軸で同じ項目を2回以上選んでいます。")
    return errors


def write_workbook(
    path: str,
    table: AnalysisTable,
    uart: UartData,
    logger: LoggerData,
    graphs: list[GraphSetting],
    progress: ProgressFn | None = None,
) -> None:
    """Excel ファイルを出力する。保存に失敗した場合は OSError（PermissionError）を送出する。"""
    if not table.rows:
        raise ValueError("出力するデータ行がありません。")
    folder = os.path.dirname(os.path.abspath(path))
    if not os.path.isdir(folder):
        raise FileNotFoundError(f"出力先フォルダがありません: {folder}")
    wb = Workbook(write_only=True)
    used: set[str] = set()

    if progress:
        progress("解析シートを書き込み中…")
    ws_a = wb.create_sheet(sanitize_sheet_name(SHEET_ANALYSIS, used))
    ws_a.sheet_properties.tabColor = TAB_COLOR_ANALYSIS
    col_of = _write_analysis(ws_a, table, uart, logger, progress)

    if graphs:
        ws_g = wb.create_sheet(sanitize_sheet_name(SHEET_GRAPH, used))
        ws_g.sheet_properties.tabColor = TAB_COLOR_GRAPH
        _add_charts(ws_g, ws_a, table, graphs, col_of)

    if progress:
        progress("UART の元データを書き込み中…")
    ws_u = wb.create_sheet(sanitize_sheet_name(uart.name, used))
    _write_uart_raw(ws_u, uart)

    if progress:
        progress("ロガーの元データを書き込み中…")
    ws_l = wb.create_sheet(sanitize_sheet_name(logger.name, used))
    _write_logger_raw(ws_l, logger, progress)

    if progress:
        progress("ファイルを保存中…")
    wb.save(path)
